from __future__ import annotations

import logging
from pathlib import Path

logger = logging.getLogger(__name__)

TEXT_SUFFIXES = {".txt", ".md", ".csv", ".json", ".xml", ".html", ".htm"}
DOCX_SUFFIXES = {".docx"}
PDF_SUFFIXES = {".pdf"}
XLSX_SUFFIXES = {".xlsx"}


class AttachmentReader:
    @staticmethod
    def read_text(file_path: Path, max_chars: int = 4000) -> str:
        suffix = file_path.suffix.lower()

        if suffix in TEXT_SUFFIXES:
            try:
                content = file_path.read_text(encoding="utf-8", errors="replace")
            except Exception as exc:
                logger.warning("Failed to read text file %s: %s", file_path.name, exc)
                return ""
        elif suffix in PDF_SUFFIXES:
            content = AttachmentReader._read_pdf(file_path)
        elif suffix in DOCX_SUFFIXES:
            content = AttachmentReader._read_docx(file_path)
        elif suffix in XLSX_SUFFIXES:
            content = AttachmentReader._read_xlsx(file_path)
        else:
            return ""

        if not content:
            return ""

        if len(content) > max_chars:
            content = content[:max_chars] + "\n\n[内容过长，已截断...]"
        return content

    @staticmethod
    def _read_pdf(file_path: Path) -> str:
        try:
            import pdfplumber
        except ImportError:
            logger.warning("pdfplumber not installed, skipping PDF: %s", file_path.name)
            return ""

        try:
            with pdfplumber.open(file_path) as pdf:
                parts = []
                for page in pdf.pages[:20]:
                    text = page.extract_text()
                    if text:
                        parts.append(text)
                return "\n".join(parts)
        except Exception as exc:
            logger.warning("Failed to read PDF %s: %s", file_path.name, exc)
            return ""

    @staticmethod
    def _read_docx(file_path: Path) -> str:
        try:
            from docx import Document
        except ImportError:
            logger.warning("python-docx not installed, skipping DOCX: %s", file_path.name)
            return ""

        try:
            doc = Document(file_path)
            return "\n".join(p.text for p in doc.paragraphs if p.text.strip())
        except Exception as exc:
            logger.warning("Failed to read DOCX %s: %s", file_path.name, exc)
            return ""

    @staticmethod
    def _read_xlsx(file_path: Path) -> str:
        try:
            import openpyxl
        except ImportError:
            logger.warning("openpyxl not installed, skipping XLSX: %s", file_path.name)
            return ""

        try:
            wb = openpyxl.load_workbook(file_path, data_only=True)
            parts = []
            for sheet_name in wb.sheetnames[:5]:
                ws = wb[sheet_name]
                parts.append(f"--- Sheet: {sheet_name} ---")
                for row in ws.iter_rows(values_only=True, max_row=200):
                    parts.append("\t".join(str(c) if c is not None else "" for c in row))
            wb.close()
            return "\n".join(parts)
        except Exception as exc:
            logger.warning("Failed to read XLSX %s: %s", file_path.name, exc)
            return ""

    def summarize(self, directory: Path, max_files: int = 5, max_chars_per_file: int = 4000) -> str:
        parts = []
        count = 0
        try:
            entries = list(directory.iterdir())
        except (FileNotFoundError, PermissionError, OSError) as exc:
            return f"(目录无法访问: {exc})"

        for f in sorted(entries):
            if not f.is_file():
                continue
            content = self.read_text(f, max_chars=max_chars_per_file)
            if content:
                parts.append(f"=== {f.name} ===\n{content}")
                count += 1
                if count >= max_files:
                    total = sum(1 for _ in entries if _.is_file())
                    parts.append(f"\n[共 {total} 个文件，已读取前 {max_files} 个]")
                    break
            else:
                parts.append(f"=== {f.name} ===\n[无法读取格式或内容为空]")
                count += 1
        return "\n\n".join(parts) if parts else "(无附件内容)"
